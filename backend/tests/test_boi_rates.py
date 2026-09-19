"""Bank of Israel rate fetching, and the appendix endpoint's choice of
rate source. The network is never touched: the fetcher is injected."""

from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.auth import hash_password
from app.db.models import ClientORM, StatementORM, UserORM
from app.db.session import SessionLocal, init_db
from app.main import app
from app.mapping.classification import classify_trade
from app.models.transactions import Currency, HoldingTerm, NormalizedStatement, SourceRef, Trade
from app.services import boi_rates
from app.services.boi_rates import BoiRatesUnavailable, fetch_boi_rates, fetch_rates_covering, parse_boi_csv

HEADER = "SERIES_CODE,FREQ,BASE_CURRENCY,COUNTER_CURRENCY,UNIT_MEASURE,DATA_TYPE,DATA_SOURCE,TIME_COLLECT,CONF_STATUS,PUB_WEBSITE,UNIT_MULT,COMMENTS,TIME_PERIOD,OBS_VALUE,RELEASE_STATUS\n"


def _csv(*rows: tuple[str, float]) -> str:
    return HEADER + "".join(f"RER_USD_ILS,D,USD,ILS,ILS,OF00,BOI_MRKT,V,F,Y,0,,{d},{v},YP\n" for d, v in rows)


def test_parse_boi_csv_reads_date_and_rate():
    rates = parse_boi_csv(_csv(("2025-03-05", 3.621), ("2025-03-06", 3.615)), Currency.USD)
    assert rates == {(Currency.USD, date(2025, 3, 5)): 3.621, (Currency.USD, date(2025, 3, 6)): 3.615}


def test_unexpected_format_is_rejected_not_guessed():
    with pytest.raises(BoiRatesUnavailable):
        parse_boi_csv("<html>maintenance</html>", Currency.USD)
    with pytest.raises(BoiRatesUnavailable):
        parse_boi_csv(_csv(("2025-03-05", -1.0)), Currency.USD)


def test_empty_answer_is_an_error():
    with pytest.raises(BoiRatesUnavailable):
        fetch_boi_rates(Currency.USD, date(2025, 1, 1), date(2025, 1, 5), fetch=lambda url: HEADER)


def test_shekel_has_no_series():
    with pytest.raises(BoiRatesUnavailable):
        fetch_boi_rates(Currency.ILS, date(2025, 1, 1), date(2025, 1, 5), fetch=lambda url: HEADER)


def test_coverage_starts_ten_days_before_the_earliest_date_for_each_currency():
    urls: list[str] = []

    def fake(url: str) -> str:
        urls.append(url)
        return _csv(("2025-03-05", 3.6))

    fetch_rates_covering({Currency.USD, Currency.ILS}, [date(2025, 3, 20), date(2025, 3, 6)], fetch=fake)
    assert len(urls) == 1  # ILS needs no rate
    assert "RER_USD_ILS" in urls[0]
    assert "startperiod=2025-02-24" in urls[0] and "endperiod=2025-03-20" in urls[0]


# ---- endpoint ----

SRC = SourceRef(statement_id="s", page=1, table_name="t", row_text="x")


def _statement_with_one_trade() -> tuple[str, str, str, str]:
    """Creates a user, client and statement; returns (email, password, statement_id, client_id)."""
    init_db()
    email, password = f"{uuid.uuid4()}@example.com", "test-password-123"
    trade = Trade(
        symbol="ABC", close_date=date(2025, 3, 10), proceeds=0, cost_basis=0, realized_pnl=100.0,
        holding_term=HoldingTerm.SHORT, source=SRC,
    )
    statement = NormalizedStatement(
        statement_id="s", period_start=date(2025, 1, 1), period_end=date(2025, 12, 31), trades=[trade]
    )
    db = SessionLocal()
    try:
        db.add(UserORM(email=email, password_hash=hash_password(password), full_name="T", role="employee"))
        client = ClientORM(full_name="test-client-" + uuid.uuid4().hex[:8])
        db.add(client)
        db.flush()
        stmt = StatementORM(
            client_id=client.id, tax_year=2025, broker="IBKR", original_filename="x.pdf",
            normalized_statement_json=statement.model_dump(mode="json"),
            classified_items_json=[classify_trade(trade).model_dump(mode="json")],
            overrides_json={},
        )
        db.add(stmt)
        db.commit()
        return email, password, stmt.id, client.id
    finally:
        db.close()


def _cleanup(statement_id: str, client_id: str) -> None:
    db = SessionLocal()
    try:
        db.query(StatementORM).filter(StatementORM.id == statement_id).delete()
        db.query(ClientORM).filter(ClientORM.id == client_id).delete()
        db.commit()
    finally:
        db.close()


def _post(client: TestClient, statement_id: str, headers: dict, body: dict):
    return client.post(f"/statements/{statement_id}/appendix.xlsx", json=body, headers=headers)


def _rate_source_cell(content: bytes) -> str | None:
    ws = load_workbook(io.BytesIO(content))["נספח ג"]
    for row in ws.iter_rows(values_only=True):
        if row[0] == "מקור שערי ההמרה":
            return row[1]
    return None


def test_endpoint_uses_bank_of_israel_rates_when_none_are_sent(monkeypatch):
    email, password, statement_id, client_id = _statement_with_one_trade()
    seen: dict = {}

    def fake_cover(currencies, dates):
        seen["currencies"], seen["dates"] = currencies, dates
        return {(Currency.USD, date(2025, 3, 10)): 3.7}

    monkeypatch.setattr("app.api.routes_reports.fetch_rates_covering", fake_cover)
    try:
        with TestClient(app) as client:
            token = client.post("/auth/login", json={"email": email, "password": password}).json()["token"]
            r = _post(client, statement_id, {"Authorization": f"Bearer {token}"}, {})
        assert r.status_code == 200, r.text
        assert seen["currencies"] == {Currency.USD} and date(2025, 3, 10) in seen["dates"]
        assert "בנק ישראל" in _rate_source_cell(r.content)
    finally:
        _cleanup(statement_id, client_id)


def test_endpoint_with_manual_rates_only_never_calls_the_bank(monkeypatch):
    email, password, statement_id, client_id = _statement_with_one_trade()

    def boom(*a, **k):
        raise AssertionError("must not fetch when only manual rates were sent")

    monkeypatch.setattr("app.api.routes_reports.fetch_rates_covering", boom)
    try:
        with TestClient(app) as client:
            token = client.post("/auth/login", json={"email": email, "password": password}).json()["token"]
            body = {"fx_rates": [{"currency": "USD", "on_date": "2025-01-01", "rate": 3.6}]}
            r = _post(client, statement_id, {"Authorization": f"Bearer {token}"}, body)
        assert r.status_code == 200, r.text
        assert _rate_source_cell(r.content) == "שערים שהוזנו ידנית"
    finally:
        _cleanup(statement_id, client_id)


def test_endpoint_reports_a_clear_error_when_the_bank_is_unreachable(monkeypatch):
    email, password, statement_id, client_id = _statement_with_one_trade()

    def down(currencies, dates):
        raise BoiRatesUnavailable("לא ניתן להתחבר לבנק ישראל")

    monkeypatch.setattr("app.api.routes_reports.fetch_rates_covering", down)
    try:
        with TestClient(app) as client:
            token = client.post("/auth/login", json={"email": email, "password": password}).json()["token"]
            r = _post(client, statement_id, {"Authorization": f"Bearer {token}"}, {})
        assert r.status_code == 422
        assert "בנק ישראל" in r.json()["detail"]
    finally:
        _cleanup(statement_id, client_id)


def test_endpoint_with_boi_off_and_no_manual_rates_is_rejected():
    email, password, statement_id, client_id = _statement_with_one_trade()
    try:
        with TestClient(app) as client:
            token = client.post("/auth/login", json={"email": email, "password": password}).json()["token"]
            r = _post(client, statement_id, {"Authorization": f"Bearer {token}"}, {"use_boi_rates": False})
        assert r.status_code == 422
    finally:
        _cleanup(statement_id, client_id)


def test_module_exposes_no_network_call_at_import_time():
    assert callable(boi_rates.fetch_boi_rates)
