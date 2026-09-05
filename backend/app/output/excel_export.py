"""Renders an AppendixReport to an .xlsx workbook the accountant can
open, check, and copy into the official Tax Authority filing."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.services.report_service import AppendixReport

_HEADER_FONT = Font(bold=True)
_RTL_ALIGN = Alignment(horizontal="right", readingOrder=2)


def _set_rtl(ws: Worksheet) -> None:
    ws.sheet_view.rightToLeft = True


def _write_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.alignment = _RTL_ALIGN


def build_appendix_workbook(report: AppendixReport) -> bytes:
    wb = Workbook()

    ws_d = wb.active
    ws_d.title = "נספח ד"
    _set_rtl(ws_d)
    _write_header(ws_d, 1, ["שדה (הכנסה)", "שדה (מס ששולם בחו\"ל)", "תיאור", "סה\"כ הכנסה (₪)", "סה\"כ מס ששולם (₪)", "מספר פריטים", "דורש אימות רו\"ח"])
    r = 2
    for t in report.nispach_d.field_totals:
        ws_d.append([t.field_income, t.field_tax_paid, t.label, t.total_income_ils, t.total_tax_paid_ils, t.item_count, "כן" if t.needs_review else "לא"])
        r += 1
    ws_d.append([])
    ws_d.append(["סה\"כ הכנסות חו\"ל (מועבר לטופס 1301, שדה 130/159)", "", "", report.nispach_d.total_foreign_income_ils])

    ws_d.append([])
    ws_d.append(["פרוטים לנספח ד' (עמוד המשלמים)"])
    _write_header(ws_d, ws_d.max_row + 1, ["שדה", "משלם", "מדינה", "סכום (₪)"])
    for row in report.nispach_d.payer_details:
        ws_d.append([row.field_income, row.payer, row.country, row.amount_ils])

    ws_c = wb.create_sheet("נספח ג")
    _set_rtl(ws_c)
    _write_header(ws_c, 1, ["מדרגת מס", "רווח הון (₪)", "מספר עסקאות"])
    for b in report.nispach_c.bracket_totals:
        ws_c.append([f"{b.tax_rate_percent}%", b.gross_gain_ils, b.item_count])
    ws_c.append([])
    ws_c.append(["סה\"כ המכירות (₪) - הערכה", report.nispach_c.total_sale_proceeds_ils])
    ws_c.append(["סה\"כ רווח הון (₪) - לפני קיזוז הפסדים", report.nispach_c.total_gain_ils])
    ws_c.append([])
    ws_c.append([
        "הערה: הטבלה כוללת רק רווח הון גולמי לפני קיזוזי הפסדים משנים קודמות/שוטפים. "
        "יש להשלים את הקיזוזים במערכת רשות המסים בהתאם לנתוני הלקוח."
    ])

    ws_e = wb.create_sheet("הסבר ומעקב")
    _set_rtl(ws_e)
    _write_header(
        ws_e,
        1,
        [
            "סוג", "מזהה מקור", "טבלת מקור בדוח", "שורת מקור",
            "סכום מקורי", "שער המרה", "שער חלופי (fallback)", "סכום בש\"ח",
            "שדה יעד", "דורש אימות", "הערת אימות",
        ],
    )
    for row in report.explanation_rows:
        ws_e.append(
            [
                row.source_kind,
                row.source_id,
                row.source_table,
                row.source_row_text,
                row.amount_source_ccy,
                row.fx_rate,
                "כן" if row.fx_rate_is_fallback else "לא",
                row.amount_ils,
                row.target_field or "",
                "כן" if row.needs_review else "לא",
                row.review_reason or "",
            ]
        )

    for ws in (ws_d, ws_c, ws_e):
        for column_cells in ws.columns:
            length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
