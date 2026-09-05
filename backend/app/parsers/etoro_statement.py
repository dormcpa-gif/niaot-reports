"""Parser for eToro's XLSX "Account Statement" export.

Grounding note (important — read before trusting this like the IBKR
parser): unlike IBKR, we do not have a real eToro export to test against.
The column headers below come from `masbug/etoro-edavki` (an actively
maintained, multi-year open-source tool that Slovenian eToro users run
against their real statements for tax filing:
https://github.com/masbug/etoro-edavki) and from `weirdapps/etoro_statement`.
Both document eToro's header text across format versions 2022-2025.2,
which is what the header-matching below tolerates. That's meaningfully
more solid than guessing, but it is still **unverified against a live
eToro file in this codebase** — treat results from this parser with
extra scrutiny until we've run it against a real export, the same way
we validated the IBKR parser against the user's real statement.

eToro-specific things worth knowing:
- All the USD-denominated columns this parser reads are always present
  regardless of the account's display currency (eToro's own EUR-focused
  tooling relies on this too), so base_currency is fixed to USD here.
- The "Dividends" sheet has clean, pre-computed net/withholding columns
  per payment -- no need to reconstruct anything.
- The "Closed Positions" sheet gives Profit(USD) per closed position
  directly, but has no reliable instrument-symbol column of its own; the
  reference tool resolves the symbol by cross-referencing the "Account
  Activity" sheet's Position ID. We don't do that cross-reference here
  (scope control) -- closed positions are labeled by Position ID, with
  the raw "Action" text kept for traceability in the explanation report.
- eToro doesn't expose a clean "gross sale proceeds" figure the way IBKR
  does, so `sale_proceeds_by_symbol` is left empty (Nispach C's "total
  sales" line will show 0 for eToro-only statements, not a guess).
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from app.models.transactions import (
    Currency,
    Dividend,
    FeeItem,
    HoldingTerm,
    InterestItem,
    NormalizedStatement,
    SourceRef,
    Trade,
)

_REQUIRED_SHEETS = {"Dividends", "Closed Positions"}

_DATE_FORMATS = (
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
)


def _normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _header_index_map(header_row: tuple) -> dict[str, int]:
    return {_normalize_header(cell): i for i, cell in enumerate(header_row) if _normalize_header(cell)}


def _find_column(headers: dict[str, int], *candidates: str) -> int | None:
    for candidate in candidates:
        if candidate in headers:
            return headers[candidate]
    return None


def _parse_cell_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_cell_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().rstrip("%")
    try:
        return float(text)
    except ValueError:
        return None


class EToroStatementParser:
    broker_name = "eToro"

    def parse(self, file_bytes: bytes, statement_id: str) -> NormalizedStatement:
        try:
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            raise ValueError(
                f"לא ניתן היה לפענח את הקובץ כ-Excel תקין (שגיאה: {e}). "
                "ודאו שזהו קובץ ה-XLSX שהתקבל מ-eToro (Account Statement)."
            ) from e

        missing = _REQUIRED_SHEETS - set(wb.sheetnames)
        if missing:
            raise ValueError(
                f"הקובץ אינו נראה כמו דוח Account Statement תקין של eToro "
                f"(חסרות הלשוניות: {', '.join(sorted(missing))}; לשוניות שנמצאו: {', '.join(wb.sheetnames)})."
            )

        dividends: list[Dividend] = []
        trades: list[Trade] = []
        all_dates: list[date] = []

        dividends_sheet = wb["Dividends"]
        rows = dividends_sheet.iter_rows(values_only=True)
        header = _header_index_map(next(rows, ()))
        col_date = _find_column(header, "Date of Payment")
        col_name = _find_column(header, "Instrument Name")
        col_net = _find_column(header, "Net Dividend Received (USD)")
        col_wht = _find_column(header, "Withholding Tax Amount (USD)")
        col_position = _find_column(header, "Position ID")
        dividends_columns_recognized = col_date is not None and col_net is not None

        if dividends_columns_recognized:
            for row in rows:
                pay_date = _parse_cell_date(row[col_date]) if col_date < len(row) else None
                net = _parse_cell_float(row[col_net]) if col_net < len(row) else None
                if pay_date is None or net is None:
                    continue
                wht = (_parse_cell_float(row[col_wht]) if col_wht is not None and col_wht < len(row) else None) or 0.0
                name = str(row[col_name]).strip() if col_name is not None and col_name < len(row) and row[col_name] else "Unknown"
                position_id = row[col_position] if col_position is not None and col_position < len(row) else ""
                all_dates.append(pay_date)
                dividends.append(
                    Dividend(
                        pay_date=pay_date,
                        symbol=name,
                        description="eToro Dividend",
                        gross_amount=round(net + wht, 2),
                        withholding_tax=round(wht, 2),
                        source=SourceRef(
                            statement_id=statement_id,
                            page=0,
                            table_name="Dividends",
                            row_text=f"{pay_date} {name} net={net} wht={wht} position={position_id}",
                        ),
                    )
                )

        positions_sheet = wb["Closed Positions"]
        rows = positions_sheet.iter_rows(values_only=True)
        header = _header_index_map(next(rows, ()))
        col_close_date = _find_column(header, "Close Date")
        col_profit = _find_column(header, "Profit(USD)")
        col_action = _find_column(header, "Action")
        col_position = _find_column(header, "Position ID")
        positions_columns_recognized = col_close_date is not None and col_profit is not None

        if positions_columns_recognized:
            for row in rows:
                close_date = _parse_cell_date(row[col_close_date]) if col_close_date < len(row) else None
                profit = _parse_cell_float(row[col_profit]) if col_profit < len(row) else None
                if close_date is None or profit is None or profit == 0:
                    continue
                action = str(row[col_action]).strip() if col_action is not None and col_action < len(row) and row[col_action] else ""
                position_id = row[col_position] if col_position is not None and col_position < len(row) else "?"
                all_dates.append(close_date)
                trades.append(
                    Trade(
                        symbol=f"Position {position_id}",
                        close_date=close_date,
                        proceeds=0.0,
                        cost_basis=0.0,
                        realized_pnl=round(profit, 2),
                        holding_term=HoldingTerm.SHORT,  # eToro doesn't expose a holding-period split; not used for Nispach C bracket selection anyway
                        source=SourceRef(
                            statement_id=statement_id,
                            page=0,
                            table_name="Closed Positions",
                            row_text=f"{close_date} position={position_id} profit={profit} action={action}",
                        ),
                    )
                )

        if not dividends_columns_recognized and not positions_columns_recognized:
            raise ValueError(
                "לא זוהו העמודות הצפויות בלשוניות Dividends ו-Closed Positions. "
                "ודאו שזהו קובץ ה-XLSX המלא (Account Statement) שהורד מ-eToro, ולא קובץ שעבר עריכה שהסירה כותרות עמודות."
            )

        if all_dates:
            period_start, period_end = min(all_dates), max(all_dates)
        else:
            # Headers matched but there were zero dividend/closed-position rows
            # (a legitimately empty period). eToro exports don't carry an
            # explicit statement-period header we can fall back to, so we
            # can't state a real period here -- this only affects the (rare)
            # all-empty case, where there's nothing to convert/report anyway.
            period_start = period_end = date.today()

        return NormalizedStatement(
            statement_id=statement_id,
            broker=self.broker_name,
            period_start=period_start,
            period_end=period_end,
            base_currency=Currency.USD,
            dividends=dividends,
            interest=[],
            trades=trades,
            fees=[],
            sale_proceeds_by_symbol={},
        )
